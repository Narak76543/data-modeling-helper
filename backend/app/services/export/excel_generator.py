import io
import re
from typing import List, Optional, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.schemas.entity import EntityCreate
from app.schemas.field import FieldCreate


# Data Type Mapping Table: Internal SQL Types -> Field-Spec Display Types
SQL_TYPE_TO_DISPLAY_MAP: Dict[str, str] = {
    "VARCHAR": "String",
    "TEXT": "Text",
    "INTEGER": "Integer",
    "BIGINT": "BigInteger",
    "BOOLEAN": "Boolean",
    "TIMESTAMP": "Timestamp",
    "DATE": "Date",
    "NUMERIC": "Decimal",
    "UUID": "UUID",
    "JSONB": "JSON",
}

# Standard Default Length / Precision per SQL Type
DEFAULT_LENGTH_MAP: Dict[str, str] = {
    "VARCHAR": "255",
    "NUMERIC": "10,2",
    "UUID": "36",
}

# Acronyms to keep uppercase in human-readable labels
KNOWN_ACRONYMS = {"id", "fk", "pk", "url", "uri", "ip", "api", "uuid", "sql", "db", "lc", "doc"}


class ExcelExportGenerator:
    """Service for generating styled Excel (.xlsx) field-specification workbooks for backend handoff."""

    @staticmethod
    def _format_field_label(name: str, explicit_label: Optional[str] = None) -> str:
        """Convert a raw snake_case or camelCase field name to a human-readable title-case label."""
        if explicit_label and explicit_label.strip():
            return explicit_label.strip()

        if not name:
            return ""

        # Normalize camelCase to snake_case first
        s1 = re.sub(r"(.)([A-Z][a-z]+)", r"\1_\2", name)
        s2 = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", s1).lower()

        tokens = s2.split("_")
        formatted_tokens = []
        for t in tokens:
            if t.lower() in KNOWN_ACRONYMS:
                formatted_tokens.append(t.upper())
            else:
                formatted_tokens.append(t.capitalize())

        return " ".join(formatted_tokens)

    @staticmethod
    def _resolve_data_type(sql_type: str) -> str:
        """Map internal SQL data types to display types."""
        normalized = (sql_type or "VARCHAR").strip().upper()
        return SQL_TYPE_TO_DISPLAY_MAP.get(normalized, normalized.capitalize())

    @staticmethod
    def _resolve_length(sql_type: str, explicit_length: Optional[str] = None) -> str:
        """Resolve field length/precision fallback."""
        if explicit_length and str(explicit_length).strip():
            return str(explicit_length).strip()

        normalized = (sql_type or "").strip().upper()
        return DEFAULT_LENGTH_MAP.get(normalized, "")

    @staticmethod
    def _resolve_man_opt(field: FieldCreate) -> str:
        """
        Placeholder convention for Man.Opt column:
        - 'M' for NOT NULL fields
        - 'PK' or 'M, PK' for Primary Keys
        - Blank for optional/nullable fields
        (Adjustable once exact team reference format is confirmed)
        """
        is_mandatory = not field.is_nullable
        is_pk = field.is_primary_key

        if is_pk and is_mandatory:
            return "M, PK"
        elif is_pk:
            return "PK"
        elif is_mandatory:
            return "M"
        return ""

    @staticmethod
    def _resolve_description(
        field: FieldCreate,
        entity_map: Dict[str, EntityCreate],
    ) -> str:
        """Generate field description fallback."""
        if field.description and field.description.strip():
            return field.description.strip()

        if field.is_foreign_key:
            target_entity = entity_map.get(field.references_entity_id or "")
            target_field_name = ""
            if target_entity and target_entity.fields:
                for tf in target_entity.fields:
                    if getattr(tf, "id", None) == field.references_field_id:
                        target_field_name = tf.name
                        break
                if not target_field_name and target_entity.fields:
                    target_field_name = target_entity.fields[0].name

            if target_entity:
                target_str = f"{target_entity.name}.{target_field_name}" if target_field_name else target_entity.name
                return f"References {target_str}"
            return "Foreign key reference"

        if field.is_primary_key:
            return "Primary key identifier"

        return ""

    @classmethod
    def generate_workbook(
        cls,
        entities: List[EntityCreate],
        project_name: str = "Data Model Specification",
    ) -> bytes:
        """
        Generate a styled Excel workbook with single stacked sheet layout.
        Returns binary content as bytes.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Field Specification"
        ws.views.sheetView[0].showGridLines = True

        # Color & Font Tokens (aligned with docs/ui-style-guide.md)
        color_accent = "1E3A5F"      # Blueprint navy header fill
        color_title_bg = "F0F4F8"    # Subtle raised entity header
        color_text_dark = "14181B"   # Ink text
        color_border = "DADDD9"      # Grid border
        color_zebra = "F9FAFB"       # Alternating row background

        font_title = Font(name="Calibri", size=11, bold=True, color=color_accent)
        font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_data = Font(name="Calibri", size=9.5, color=color_text_dark)
        font_mono = Font(name="Consolas", size=9, color=color_text_dark)

        fill_title = PatternFill(start_color=color_title_bg, end_color=color_title_bg, fill_type="solid")
        fill_header = PatternFill(start_color=color_accent, end_color=color_accent, fill_type="solid")
        fill_zebra = PatternFill(start_color=color_zebra, end_color=color_zebra, fill_type="solid")
        fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        thin_border = Border(
            left=Side(style="thin", color=color_border),
            right=Side(style="thin", color=color_border),
            top=Side(style="thin", color=color_border),
            bottom=Side(style="thin", color=color_border),
        )

        header_border = Border(
            left=Side(style="thin", color=color_accent),
            right=Side(style="thin", color=color_accent),
            top=Side(style="thin", color=color_accent),
            bottom=Side(style="thin", color=color_accent),
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        headers = [
            "No",
            "Field Label",
            "Field Name",
            "Data Type",
            "Length",
            "Man.Opt",
            "LOV",
            "Default Value",
            "Description",
        ]

        # Entity lookup map for foreign key reference resolution
        entity_map: Dict[str, EntityCreate] = {
            getattr(e, "id", str(idx)): e for idx, e in enumerate(entities)
        }

        current_row = 1

        # Render each entity as a structured table block
        for entity_idx, entity in enumerate(entities):
            fields = entity.fields or []

            # 1. Table Title Row (Merged A to I)
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
            title_cell = ws.cell(row=current_row, column=1)
            title_cell.value = f"Table: {entity.name}"
            title_cell.font = font_title
            title_cell.fill = fill_title
            title_cell.alignment = align_left
            ws.row_dimensions[current_row].height = 24

            for col_idx in range(1, 10):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.fill = fill_title
                cell.border = thin_border

            current_row += 1

            # 2. Column Headers Row
            ws.row_dimensions[current_row].height = 22
            for col_idx, h_text in enumerate(headers, start=1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.value = h_text
                cell.font = font_header
                cell.fill = fill_header
                cell.border = header_border
                cell.alignment = align_center if h_text in {"No", "Data Type", "Length", "Man.Opt", "LOV"} else align_left

            current_row += 1

            # 3. Data Rows per Field
            if not fields:
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
                empty_cell = ws.cell(row=current_row, column=1)
                empty_cell.value = "(No fields defined)"
                empty_cell.font = Font(name="Calibri", size=9.5, italic=True, color="8A9195")
                empty_cell.alignment = align_center
                for col_idx in range(1, 10):
                    ws.cell(row=current_row, column=col_idx).border = thin_border
                ws.row_dimensions[current_row].height = 20
                current_row += 1
            else:
                for field_idx, field in enumerate(fields, start=1):
                    ws.row_dimensions[current_row].height = 20
                    row_fill = fill_zebra if field_idx % 2 == 0 else fill_white

                    field_label = cls._format_field_label(field.name, getattr(field, "label", None))
                    field_name = field.name
                    data_type = cls._resolve_data_type(field.data_type)
                    length_val = cls._resolve_length(field.data_type, getattr(field, "length", None))
                    man_opt_val = cls._resolve_man_opt(field)
                    # LOV (List of Values): intentionally unbuilt pending reference use case
                    lov_val = ""
                    default_val = field.default_value or ""
                    description_val = cls._resolve_description(field, entity_map)

                    row_values = [
                        (field_idx, align_center, font_data),
                        (field_label, align_left, font_data),
                        (field_name, align_left, font_mono),
                        (data_type, align_center, font_data),
                        (length_val, align_center, font_data),
                        (man_opt_val, align_center, font_data),
                        (lov_val, align_center, font_data),
                        (default_val, align_left, font_mono if default_val else font_data),
                        (description_val, align_left, font_data),
                    ]

                    for col_idx, (val, alignment, font_style) in enumerate(row_values, start=1):
                        cell = ws.cell(row=current_row, column=col_idx)
                        cell.value = val
                        cell.font = font_style
                        cell.alignment = alignment
                        cell.fill = row_fill
                        cell.border = thin_border

                    current_row += 1

            # 4. Spacing: 2 empty rows between entity tables
            current_row += 2

        # Auto-adjust column widths with sensible minimums
        min_widths = {
            1: 6,   # No
            2: 20,  # Field Label
            3: 20,  # Field Name
            4: 14,  # Data Type
            5: 10,  # Length
            6: 12,  # Man.Opt
            7: 10,  # LOV
            8: 16,  # Default Value
            9: 34,  # Description
        }

        for col_idx in range(1, 10):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for row_idx in range(1, current_row):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val:
                    max_len = max(max_len, len(str(val)))
            base_min = min_widths.get(col_idx, 12)
            ws.column_dimensions[col_letter].width = max(base_min, min(max_len + 3, 50))

        # Save workbook to memory buffer
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return bio.getvalue()
